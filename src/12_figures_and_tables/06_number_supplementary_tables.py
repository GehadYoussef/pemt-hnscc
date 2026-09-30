"""Give the supplementary tables their published numbers.

The table builders write each sheet under a working name (for example S22_pemt_gene_partition).
config/supplementary_tables.tsv maps every working table id to its number in the paper, where
tables are numbered in order of first citation. The sheets and their Index rows are renamed and put
in published order, and the table cross-references in the Index and in each caption row are
rewritten.

Reads and rewrites: results/tables/Supplementary_Data_1.xlsx
Run after 05_extend_supplementary_tables.py:
    python src/12_figures_and_tables/06_number_supplementary_tables.py
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "results" / "tables" / "Supplementary_Data_1.xlsx"
MAP = ROOT / "config" / "supplementary_tables.tsv"

LETTERS = "bcdefghijklmnopqrstuvwxyz"   # the first sheet of a table is unlettered
XREF = re.compile(r"(?<![A-Za-z0-9:.])S(\d+)([a-z]?)(?=[_\s.,)]|$)")


def parse(tok: str) -> tuple[int, str]:
    m = re.fullmatch(r"S(\d+)([a-z]?)", tok)
    return int(m.group(1)), m.group(2)


def sheet_id(name: str) -> tuple[int, str] | None:
    m = re.match(r"S(\d+)([a-z]?)_", name)
    return (int(m.group(1)), m.group(2)) if m else None


def order_key(name: str):
    sid = sheet_id(name)
    return (sid[0], LETTERS.index(sid[1]) + 1 if sid[1] else 0, name) if sid else (0, 0, "")


def rename(name: str, mapping) -> str:
    """Published sheet name, cut to the 31 characters Excel allows."""
    n, l = mapping[sheet_id(name)]
    return (f"S{n}{l}_" + name.split("_", 1)[1])[:31]


def xref(text: str, mapping) -> str:
    def sub(m):
        key = (int(m.group(1)), m.group(2))
        if key not in mapping:
            return m.group(0)
        n, l = mapping[key]
        return f"S{n}{l}"
    return XREF.sub(sub, text)


def main() -> None:
    table = pd.read_csv(MAP, sep="\t")
    mapping = {parse(w): parse(p) for w, p in zip(table["working_id"], table["published_id"])}
    wb = openpyxl.load_workbook(XLSX)
    missing = [s for s in wb.sheetnames if sheet_id(s) and sheet_id(s) not in mapping]
    if missing:
        raise SystemExit(f"sheets without a published number in {MAP.name}: {missing}")

    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(min_row=1, max_row=1):
            for c in row:
                if isinstance(c.value, str):
                    c.value = xref(c.value, mapping)
    idx = wb["Index"]
    for row in idx.iter_rows(min_row=2):
        if isinstance(row[0].value, str) and sheet_id(row[0].value):
            row[0].value = rename(row[0].value, mapping)
        if isinstance(row[1].value, str):
            row[1].value = xref(row[1].value, mapping)

    new_names = [rename(ws.title, mapping) if sheet_id(ws.title) else ws.title for ws in wb.worksheets]
    if len(set(new_names)) != len(new_names):
        raise SystemExit("renaming would create duplicate sheet names")
    for k, ws in enumerate(wb.worksheets[1:], 1):
        ws.title = f"tmp{k}"
    for ws, name in zip(wb.worksheets, new_names):
        ws.title = name

    wb._sheets = [idx] + sorted(wb.worksheets[1:], key=lambda ws: order_key(ws.title))
    body = [[c.value for c in r] for r in idx.iter_rows(min_row=2)]
    body.sort(key=lambda r: order_key(r[0]) if isinstance(r[0], str) else (10**6,))
    for i, r in enumerate(body, 2):
        for j, v in enumerate(r, 1):
            idx.cell(row=i, column=j, value=v)
    wb.save(XLSX)
    n_tables = len({n for n, _ in mapping.values()})
    print(f"{len(wb.sheetnames) - 1} sheets numbered as Supplementary Tables S1 to S{n_tables}: {XLSX}")


if __name__ == "__main__":
    main()
